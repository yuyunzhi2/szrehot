# -*- coding: utf-8 -*-
"""szrehot 每日链执行器：阶段 1 京东法拍采集 → 阶段 2 深圳法拍资产日报。

阶段 3（深圳资产日报：新闻六章节 + 情绪指数 + 租赁 + 法拍热点）需主模型按
SKILL.md 五阶段语义产出，本脚本只负责前两个确定性阶段，跑完交回主模型。

用法（工作目录 = szrehot 技能目录）：
  python -X utf8 scripts/run_daily_chain.py                 # 今天，全链
  python -X utf8 scripts/run_daily_chain.py --dry-run       # 只抓不写 xlsx，日报落临时目录
  python -X utf8 scripts/run_daily_chain.py --date 2026-10-27
  python -X utf8 scripts/run_daily_chain.py --skip-jd       # 只重出法拍日报
  python -X utf8 scripts/run_daily_chain.py --force         # 今日块已存在也重跑阶段 1

幂等：阶段 1 目标 sheet 已有当日日期行时默认跳过（原 jd-paipai 每次写日期行，
日内重复执行会留下空块），--force 强制重跑。
"""
import argparse
import datetime as dt
import os
import re
import subprocess
import sys
import tempfile

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(BASE, 'scripts')
DEFAULT_XLSX = r'D:\Desktop\法拍数据.xlsx'
DEFAULT_FAPAI_OUT = r'D:\Administrator\Documents\大瑀MD'

# 阶段 1：(采集脚本, 目标 sheet, 轨道名)
JD_TRACKS = [
    ('fetch_active.py', '京东API', '在售（进行中+预告中）'),
    ('fetch_closed.py', '京东成交API', '成交（已结束窗口内）'),
]


def parse_date(text):
    """'2026-10-28' / '2026年10月28日' / '2026/10/28' → date；None 视为今天"""
    if not text:
        return dt.date.today()
    m = re.search(r'(\d{4})\s*[-/.年]\s*(\d{1,2})\s*[-/.月]\s*(\d{1,2})', text)
    if not m:
        raise SystemExit(f'无法识别日期：{text}')
    return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def block_state(xlsx, sheet, day):
    """返回 (是否有当日日期行, 块内数据行数)。文件/sheet 缺失按无块处理。"""
    if not os.path.exists(xlsx):
        return False, 0
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return False, 0
    forms = {day.strftime('%Y年%m月%d日'), day.isoformat(), day.strftime('%Y/%m/%d')}
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    for i, row in enumerate(rows):
        cell = row[0] if row else None
        if isinstance(cell, str) and cell.strip() in forms:
            body = 0
            for nxt in rows[i + 1:]:
                v = nxt[0] if nxt else None
                if isinstance(v, str) and v.strip() in forms:
                    break
                if isinstance(v, str) and v.strip():
                    body += 1
            return True, body
    return False, 0


def run(cmd, label):
    print(f'[chain] {label}: {" ".join(cmd)}', flush=True)
    code = subprocess.run(cmd).returncode
    if code != 0:
        print(f'[chain] {label} 失败（退出码 {code}）', file=sys.stderr, flush=True)
    return code


def stage_jd(args, day):
    print(f'[chain] 阶段 1/3 京东法拍采集 → {args.xlsx}', flush=True)
    failed = []
    for script, sheet, name in JD_TRACKS:
        exists, body = block_state(args.xlsx, sheet, day)
        if exists and not args.force and not args.dry_run:
            print(f'[chain] 跳过 {sheet}（{name}）：已有 {day} 块，数据行 {body}；'
                  f'需重跑加 --force', flush=True)
            continue
        cmd = [sys.executable, '-X', 'utf8', os.path.join(SCRIPTS, 'jd_paipai', script),
               '--xlsx', args.xlsx, '--date', day.isoformat()]
        if script == 'fetch_closed.py':
            # 结束时间窗口随业务日期滚动，避免 --date 指定往日时窗口仍按今天算
            cmd += ['--since', (day - dt.timedelta(days=10)).isoformat(),
                    '--until', day.isoformat()]
        if args.dry_run:
            cmd.append('--dry-run')
        if run(cmd, f'1 京东法拍 {sheet}') != 0:
            failed.append(sheet)
    return failed


def stage_fapai(args, day):
    print('[chain] 阶段 2/3 深圳法拍资产日报', flush=True)
    out_dir = args.fapai_out_dir
    if args.dry_run:
        out_dir = os.path.join(tempfile.gettempdir(), 'szrehot_chain_dryrun')
        print(f'[chain] dry-run：日报落到临时目录 {out_dir}', flush=True)
    cmd = [sys.executable, '-X', 'utf8', os.path.join(SCRIPTS, 'generate_fapai_daily.py'),
           '--date', day.isoformat(), '--xlsx', args.xlsx, '--out-dir', out_dir]
    return run(cmd, '2 法拍资产日报')


def main():
    ap = argparse.ArgumentParser(description='szrehot 每日链：京东法拍采集 → 法拍资产日报')
    ap.add_argument('--date', default=None, help='业务日期 YYYY-MM-DD（默认今天）')
    ap.add_argument('--xlsx', default=DEFAULT_XLSX, help='法拍数据源 xlsx')
    ap.add_argument('--fapai-out-dir', default=DEFAULT_FAPAI_OUT, help='法拍日报输出目录')
    ap.add_argument('--skip-jd', action='store_true', help='跳过阶段 1（京东法拍采集）')
    ap.add_argument('--skip-fapai', action='store_true', help='跳过阶段 2（法拍资产日报）')
    ap.add_argument('--force', action='store_true', help='当日块已存在也重跑阶段 1')
    ap.add_argument('--dry-run', action='store_true', help='阶段 1 只抓不写，阶段 2 落临时目录')
    args = ap.parse_args()

    day = parse_date(args.date)
    print(f'[chain] 业务日期 {day} · 数据源 {args.xlsx}', flush=True)

    failed = [] if args.skip_jd else stage_jd(args, day)
    if not args.skip_fapai and stage_fapai(args, day) != 0:
        failed.append('法拍资产日报')

    if failed:
        print(f'[chain] 未完成：{"、".join(failed)}；其余阶段产物仍可用，'
              f'主模型按 SKILL.md 继续阶段 3 并在日报注明限制', file=sys.stderr, flush=True)
        return 1
    print('[chain] 阶段 1-2 完成。阶段 3（深圳资产日报）由主模型按 SKILL.md 五阶段执行。',
          flush=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
