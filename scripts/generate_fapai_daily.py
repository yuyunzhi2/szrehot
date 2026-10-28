# -*- coding: utf-8 -*-
"""深圳法拍资产日报生成器

从 法拍数据.xlsx 的 京东 / 阿里 / 京东成交 / 阿里成交 四个工作表中，
按日期块提取当日数据，输出 markdown 日报。

日期标记兼容三种形态：
  1) 中文文本 "2026年10月27日"（允许首尾空白）
  2) ISO 文本 "2026-10-27" / "2026/10/27" / "2026.10.27"
  3) Excel 日期序列值（如 46322）或 datetime/date 对象

呈现规则：
  - 拍卖开始时间只保留到日期（如 "11月23日 10:00:00开始" -> "11月23日"）
  - 项目名称清洗（删 省市前缀/关于/拍卖/位于/权证号等注记括号）
  - 含"二拍"的名称把"二拍"提到首位，全部二拍项独立成组置顶，组内按
    "二拍"后的城区名称首字拼音降序；其余项按城区降序；无城区的行排
    组内最后并保持原序。
  - 入选规则：
      关注（新增关注top10）：当前价与评估价均有数据，按热度降序取前 10
      成交（成交top10）：当前价与评估价均有数据；折价率 > 100% 的一律
        入选，其余按折价率由低到高补足至 10 条

用法（工作目录 = szrehot 技能目录）：
  python -X utf8 scripts/generate_fapai_daily.py                    # 今天，缺省路径
  python -X utf8 scripts/generate_fapai_daily.py --date 2026-10-27
  python -X utf8 scripts/generate_fapai_daily.py --xlsx PATH --out-dir DIR

本脚本是 szrehot 每日链的第 2 步（第 1 步 = scripts/jd_paipai/ 京东法拍采集，
第 3 步 = 本技能原有五阶段新闻管道），规则详表见 references/fapai-daily.md。
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError:
    sys.stderr.write("缺少 openpyxl，请先安装：python -m pip install openpyxl\n")
    raise

try:
    from pypinyin import lazy_pinyin
    _HAS_PINYIN = True
except ImportError:
    _HAS_PINYIN = False

DEFAULT_XLSX = r"D:\Desktop\法拍数据.xlsx"
DEFAULT_OUT_DIR = r"D:\Administrator\Documents\大瑀MD"

# 表配置：(sheet 名, 日期列 0-based, 模式)
#   watch = 新增关注：按 F 列热度降序取前 10
#   deal_jd = 京东成交：F 列含"出价"字样
#   deal_ali = 阿里成交：F 列出价人数 > 0
SHEETS = [
    ("阿里", 0, "watch"),
    ("京东", 1, "watch"),
    ("阿里成交", 0, "deal_ali"),
    ("京东成交", 1, "deal_jd"),
]

TITLE_MAP = {
    "阿里": "阿里法拍新增关注top10",
    "京东": "京东法拍新增关注top10",
    "阿里成交": "阿里法拍成交top10",
    "京东成交": "京东法拍成交top10",
}

# 概览里的简称
OVERVIEW_MAP = {
    "阿里": "阿里新增关注",
    "京东": "京东新增关注",
    "阿里成交": "阿里成交",
    "京东成交": "京东成交",
}

TOP_N = 10

DATE_PATTERNS = [
    re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"),
    re.compile(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})"),
]
PURE_NUM = re.compile(r"^\d+(?:\.\d+)?$")
HEAT_W = re.compile(r"^(\d+(?:\.\d+)?)\s*[wW万]\s*人?围观?$")
HEAT_PLAIN = re.compile(r"^(\d+(?:\.\d+)?)\s*人围观?$")

# 拍卖开始时间：截到日期
START_CJK_FULL = re.compile(r"\d{4}年\d{1,2}月\d{1,2}日")
START_CJK_MD = re.compile(r"\d{1,2}月\d{1,2}日")
START_ISO = re.compile(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}")

# 项目名称清洗
JUNK_TOKENS = ["广东省", "深圳市", "关于", "拍卖", "位于"]
JUNK_BRACKET_KEYS = ("不动产权证号", "不含室内物品", "不包含动产", "装修按现状", "带租")
BRACKET_PATS = [re.compile(r"（[^（）]*）"), re.compile(r"\([^()]*\)"),
                re.compile(r"\[[^\]]*\]"), re.compile(r"【[^】]*】")]
DUP_PREFIX = [("变卖变卖", "变卖"), ("拍卖拍卖", "拍卖"), ("二拍二拍", "二拍")]

# 城区提取
KNOWN_DISTRICTS = ["罗湖", "福田", "南山", "盐田", "宝安", "龙岗", "龙华",
                   "坪山", "光明", "大鹏", "深汕", "前海"]
LOCALITY_GENERIC = re.compile(r"([\u4e00-\u9fa5]{2})(?:新区|街道|镇|区|县)")
GENERIC_BAD = ("小区", "工业区", "商业区", "住宅区", "宿舍区", "生活区", "科技区",
               "开发区", "经济区", "保税区", "加工区", "厂区", "校区", "园区")


# ---------- 解析工具 ----------

def parse_date_cell(v):
    """单元格 -> datetime.date | None（兼容中文/ISO 文本、Excel 序列值、日期对象）"""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # Excel 序列值：1900 前后至 2100 年约合 1..80000，热度/价格列不在扫描范围
        if 20000 <= float(v) <= 80000:
            try:
                d = from_excel(v)
                return d.date() if isinstance(d, dt.datetime) else d
            except Exception:
                return None
        return None
    if isinstance(v, str):
        s = v.strip().strip("　")
        for pat in DATE_PATTERNS:
            m = pat.search(s)
            if m:
                y, mo, dy = int(m.group(1)), int(m.group(2)), int(m.group(3))
                try:
                    return dt.date(y, mo, dy)
                except ValueError:
                    return None
    return None


def is_marker(row, date_col):
    """日期标记行：日期列可解析成日期，且整行不含 URL（数据行必有网址列）"""
    if date_col >= len(row):
        return None
    d = parse_date_cell(row[date_col])
    if d is None:
        return None
    for cell in row:
        if isinstance(cell, str) and cell.strip().lower().startswith(("http://", "https://")):
            return None
    return d


def is_data_row(row):
    """数据行：项目名称（A 列）非空"""
    v = row[0] if row else None
    return isinstance(v, str) and v.strip() != ""


def parse_price_wan(v):
    """价格 -> 万元 float。兼容 284.49万 / ¥355.62万 / 2.5亿 / 裸数字(按万) / 1234元"""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if not isinstance(v, str):
        return None
    s = v.strip().replace(",", "").replace("，", "").replace("￥", "").replace("¥", "").replace(" ", "")
    if not s:
        return None
    mult = None
    if s.endswith("亿元"):
        s, mult = s[:-2], 10000.0
    elif s.endswith("亿"):
        s, mult = s[:-1], 10000.0
    elif s.endswith("万元"):
        s, mult = s[:-2], 1.0
    elif s.endswith("万"):
        s, mult = s[:-1], 1.0
    elif s.endswith("元"):
        s, mult = s[:-1], 0.0001
    try:
        num = float(s)
    except ValueError:
        return None
    return num * (mult if mult is not None else 1.0)


def parse_heat(v):
    """F 列热度 -> float | None。忽略 单价（元/平）等无关数据"""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if not isinstance(v, str):
        return None
    s = v.strip().replace(",", "").replace("，", "").replace(" ", "")
    if not s:
        return None
    if PURE_NUM.match(s):
        return float(s)
    m = HEAT_W.match(s)  # 3.63w人围观 / 3.63万围观
    if m:
        return float(m.group(1)) * 10000.0
    m = HEAT_PLAIN.match(s)  # 1900人围观
    if m:
        return float(m.group(1))
    return None  # 元/平、商业 等噪声


def fmt_wan(v):
    if v is None:
        return "—"
    s = f"{v:.2f}".rstrip("0").rstrip(".")
    return s if s else "0"


def fmt_ratio(cur, ev):
    if cur is None or ev is None or ev == 0:
        return "—"
    return f"{cur / ev * 100:.1f}%"


def fmt_start(v):
    """拍卖开始时间只保留到日期：'11月23日 10:00:00开始' -> '11月23日'"""
    if v is None:
        return "—"
    s = str(v).strip()
    if not s:
        return "—"
    for pat in (START_CJK_FULL, START_CJK_MD, START_ISO):
        m = pat.search(s)
        if m:
            return m.group(0)
    return s  # 进行中/预告中/已结束 等无日期文本


def clean_name(name):
    """项目名称清洗：删省市前缀、关于/拍卖/位于等填充词、权证号等注记括号"""
    s = name.strip()

    def _bracket_repl(m):
        seg = m.group(0)
        if any(k in seg for k in JUNK_BRACKET_KEYS):
            return ""
        return seg

    for pat in BRACKET_PATS:
        s = pat.sub(_bracket_repl, s)
    for t in JUNK_TOKENS:
        s = s.replace(t, "")
    for a, b in DUP_PREFIX:
        s = s.replace(a, b)
    s = re.sub(r"\s+", " ", s).strip()
    s = s.strip(" 、，,。.:：;；")  # 清除删词后残留的首尾标点
    return s or name.strip()


def promote_er(name):
    """含"二拍"的名称把"二拍"两字提到首位（已在首位则不变）"""
    if "二拍" not in name or name.startswith("二拍"):
        return name
    i = name.find("二拍")
    s = "二拍" + name[:i] + name[i + 2:]
    for empty in ("（）", "()", "【】", "[]"):
        s = s.replace(empty, "")
    return s


def extract_locality(name):
    """取城区名称：'二拍'名称看'二拍'之后；优先已知城区，其次通用 XX区/镇/街道"""
    seg = name.split("二拍", 1)[1] if "二拍" in name else name
    best, best_pos = None, None
    for d in KNOWN_DISTRICTS:
        p = seg.find(d)
        if p >= 0 and (best_pos is None or p < best_pos):
            best, best_pos = d, p
    if best:
        return best
    for m in LOCALITY_GENERIC.finditer(seg):
        tok = m.group(0)
        if any(tok.endswith(b) or tok == b for b in GENERIC_BAD):
            continue
        return tok
    return None


def locality_pinyin(locality):
    """城区名称首字拼音（如 '罗湖区' -> 'luo'）；无 pypinyin 时退化为字符本身"""
    if not locality:
        return ""
    ch = locality[0]
    if _HAS_PINYIN:
        try:
            py = lazy_pinyin(ch)[0]
            return py.lower() if isinstance(py, str) else ""
        except Exception:
            return ""
    return ch


def sort_by_locality(items):
    """二拍项独立成组置顶，组内按城区首字拼音降序；其余按城区降序；无城区排组内最后"""
    if not _HAS_PINYIN:
        sys.stderr.write("提示：未安装 pypinyin（python -m pip install pypinyin），"
                         "城区排序退化为字符码降序\n")

        def order(group):
            known = sorted([i for i in group if i["loc"]],
                           key=lambda x: x["loc"][0], reverse=True)
            return known + [i for i in group if not i["loc"]]
    else:
        def order(group):
            known = sorted([i for i in group if i["py"]],
                           key=lambda x: x["py"], reverse=True)
            return known + [i for i in group if not i["py"]]
    erpai = [i for i in items if i.get("is_er")]
    normal = [i for i in items if not i.get("is_er")]
    return order(erpai) + order(normal)


def md_cell(s):
    if s is None:
        return "—"
    return str(s).replace("|", "\\|").replace("\n", " ").replace("\r", " ").strip() or "—"


# ---------- 数据提取 ----------

def extract_block(rows, date_col, target):
    """定位 target 日期块；当天无有效块则回退到最新的有效块。
    返回 (data_rows, used_date, fallback: bool, block_size: int)"""
    markers = []  # (row_idx_0based, date)
    for i, r in enumerate(rows):
        d = is_marker(r, date_col)
        if d is not None:
            markers.append((i, d))
    if not markers:
        return [], None, False, 0

    blocks = []  # (date, row_idx_of_marker, data_rows)
    for k, (mi, d) in enumerate(markers):
        end = markers[k + 1][0] if k + 1 < len(markers) else len(rows)
        body = [r for r in rows[mi + 1:end] if is_data_row(r)]
        blocks.append((d, mi, body))

    for d, _, body in blocks:
        if d == target and body:
            return body, d, False, len(body)
    # 回退：取日期最新且非空的块
    valid = [b for b in blocks if b[2]]
    if not valid:
        return [], None, False, 0
    d, _, body = max(valid, key=lambda b: (b[0], b[1]))
    return body, d, True, len(body)


def make_item(r):
    """行 -> 输出项（含清洗名、二拍标记、城区与拼音键）"""
    cur = parse_price_wan(r[2]) if len(r) > 2 else None
    ev = parse_price_wan(r[3]) if len(r) > 3 else None
    start = fmt_start(r[4]) if len(r) > 4 else "—"
    name = promote_er(clean_name(str(r[0])))
    is_er = "二拍" in name
    loc = extract_locality(name)
    return {"name": md_cell(name), "cur": cur, "ev": ev,
            "start": md_cell(start), "loc": loc, "py": locality_pinyin(loc),
            "is_er": is_er}


def has_prices(it):
    """入选门槛：当前价与评估价都有有效数据（不编数）"""
    return it["cur"] is not None and it["ev"] is not None and it["ev"] > 0


def ratio_of(it):
    """折价率 = 当前价/评估价（小数），无法计算返回 None"""
    if not has_prices(it):
        return None
    return it["cur"] / it["ev"]


def build_watch_rows(body):
    """新增关注表：当前价与评估价齐全的有效热度行，按热度降序取前 10，再按城区排序"""
    items = []
    for r in body:
        heat = parse_heat(r[5]) if len(r) > 5 else None
        if heat is None:
            continue
        it = make_item(r)
        if not has_prices(it):
            continue  # 价格不全的项不入选
        it["heat"] = heat
        items.append(it)
    items.sort(key=lambda x: x["heat"], reverse=True)
    return sort_by_locality(items[:TOP_N]), len(items)


def build_deal_rows(body, mode):
    """成交表：京东按 F 含'出价'；阿里按 F 出价人数>0。

    入选：当前价与评估价齐全；折价率 > 100% 的一律入选，
    其余按折价率由低到高补足至 TOP_N 条；展示顺序仍按城区规则。
    """
    items = []
    for r in body:
        f = r[5] if len(r) > 5 else None
        if mode == "deal_jd":
            if not (isinstance(f, str) and "出价" in f):
                continue  # 忽略 'X人围观' 等
        else:  # deal_ali
            n = parse_heat(f)
            if n is None or n == 0:
                continue  # 忽略 0 与噪声
        it = make_item(r)
        if not has_prices(it):
            continue  # 价格不全的项不入选
        it["ratio"] = ratio_of(it)
        items.append(it)

    # 折价率 > 100% 必选；其余按折价率由低到高（相同则保持表内原序）
    hot = [it for it in items if it["ratio"] > 1.0]
    rest = sorted([it for it in items if it["ratio"] <= 1.0],
                  key=lambda x: x["ratio"])
    picked = hot + rest[:max(0, TOP_N - len(hot))]
    return sort_by_locality(picked), len(items)


# ---------- 输出 ----------

def render_watch_table(items):
    lines = ["| 项目名称 | 当前价（万） | 评估价（万） | 折价率 | 拍卖开始时间 |",
             "| --- | --- | --- | --- | --- |"]
    for it in items:
        lines.append(f"| {it['name']} | {fmt_wan(it['cur'])} | {fmt_wan(it['ev'])} "
                     f"| {fmt_ratio(it['cur'], it['ev'])} | {it['start']} |")
    return "\n".join(lines)


def render_deal_table(items):
    lines = ["| 项目名称 | 成交价（万） | 评估价（万） | 折价率 |",
             "| --- | --- | --- | --- |"]
    for it in items:
        lines.append(f"| {it['name']} | {fmt_wan(it['cur'])} | {fmt_wan(it['ev'])} "
                     f"| {fmt_ratio(it['cur'], it['ev'])} |")
    return "\n".join(lines)


def date_note(used, target, fallback):
    if used is None:
        return "_无数据_"
    if fallback:
        return (f"_数据日期：{used.isoformat()}（{target.isoformat()} 无有效数据块，回退取最新块）_")
    return f"_数据日期：{used.isoformat()}_"


def main():
    ap = argparse.ArgumentParser(description="生成深圳法拍资产日报")
    ap.add_argument("--date", help="运行日期，默认今天，如 2026-10-27 或 2026年10月27日")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX, help="数据源 xlsx 路径")
    ap.add_argument("--out-dir", default=DEFAULT_OUT_DIR, help="日报输出目录")
    ap.add_argument("--out", help="显式指定输出文件路径（覆盖 out-dir）")
    args = ap.parse_args()

    target = parse_date_cell(args.date) if args.date else dt.date.today()
    if target is None:
        sys.stderr.write(f"无法识别日期：{args.date}\n")
        sys.exit(2)

    xlsx = Path(args.xlsx)
    if not xlsx.exists():
        sys.stderr.write(f"数据源不存在：{xlsx}\n")
        sys.exit(2)

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    results = {}   # sheet -> dict
    order = []
    for name, date_col, mode in SHEETS:
        if name not in wb.sheetnames:
            print(f"[{name}] 工作表不存在，跳过")
            results[name] = {"items": [], "used": None, "fallback": False, "valid": 0, "block": 0, "mode": mode}
            order.append(name)
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        body, used, fallback, block_size = extract_block(rows, date_col, target)
        if mode == "watch":
            items, valid = build_watch_rows(body)
        else:
            items, valid = build_deal_rows(body, mode)
        results[name] = {"items": items, "used": used, "fallback": fallback,
                         "valid": valid, "block": block_size, "mode": mode}
        tag = "回退" if fallback else "当日"
        print(f"[{name}] {tag} 数据日期={used} 块内行数={block_size} 候选行={valid} 入选={len(items)}")
        order.append(name)
    wb.close()

    # 组装 markdown
    out_lines = [f"# 深圳法拍资产日报（{target.isoformat()}）", ""]
    overview = []
    for name in order:
        r = results[name]
        label = OVERVIEW_MAP[name]
        overview.append(f"{label} {len(r['items'])} 条" + ("（回退）" if r["fallback"] else ""))
    out_lines.append("概览：" + "｜".join(overview))
    out_lines.append("")
    out_lines.append("---")
    out_lines.append("")

    for name in order:
        r = results[name]
        out_lines.append(f"## {TITLE_MAP[name]}")
        out_lines.append("")
        out_lines.append(date_note(r["used"], target, r["fallback"]))
        out_lines.append("")
        if not r["items"]:
            out_lines.append("_无数据_")
        elif r["mode"] == "watch":
            out_lines.append(render_watch_table(r["items"]))
        else:
            out_lines.append(render_deal_table(r["items"]))
        out_lines.append("")

    out_lines.append("---")
    out_lines.append("")
    out_lines.append("注：折价率 = 价格 ÷ 评估价；「—」表示源数据缺失。价格统一折算为万元。"
                     "「二拍」项独立成组置顶、名称前置「二拍」，组内按二拍后城区首字拼音降序；"
                     "其余按城区降序。仅收录当前价与评估价均有数据的项目；"
                     "新增关注按热度取前 10；成交折价率超 100% 者必选，其余按折价率由低到高补足 10 条。")
    out_lines.append("")

    if args.out:
        out_path = Path(args.out)
    else:
        out_dir = Path(args.out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"深圳法拍资产日报-{target.isoformat()}.md"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(out_lines), encoding="utf-8")
    print(f"已生成：{out_path}")


if __name__ == "__main__":
    main()
