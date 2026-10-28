# -*- coding: utf-8 -*-
"""jd-paipai 写入层（szrehot 每日链步骤 1）：日期行 + 块内去重 + 块内排序 + 追加写入。

块格式（与“阿里成交”现有块一致）：日期行（A 列 `YYYY年MM月DD日`，单格）
下一行起为数据行（16 列，无块内表头，列定义见 references/jd-paipai.md）。
去重/排序只作用于本次写入块，不动历史块。
"""
import datetime
import sys

NAME_I, PRICE_I = 2, 4      # 物业名称 / 起拍价(万元)
WATCH_I, BID_I, START_I = 11, 10, 12  # 围观数 / 出价次数 / 拍卖开始时间


def _key(row):
    name = str(row[NAME_I] or '').strip()
    try:
        price = round(float(row[PRICE_I]), 2) if row[PRICE_I] not in (None, '') else None
    except Exception:
        price = str(row[PRICE_I])
    return (name, price)


def dedupe(rows):
    seen, out = set(), []
    for r in rows:
        k = _key(r)
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    return out, len(rows) - len(out)


def history_keys(ws):
    """本 sheet 历史键集合：(物业名称去空格, 起拍价万元2位)，与 _key 同口径。
    空名行（日期行、空行）跳过；列数不足的旧版式行跳过。"""
    keys = set()
    for row in ws.iter_rows(values_only=True):
        if len(row) <= PRICE_I:
            continue
        k = _key(row)
        if not k[0]:
            continue
        keys.add(k)
    return keys


def dedupe_history(rows, seen):
    """剔除与历史重复的行（key 已存在），返回 (保留rows, 剔除数)。"""
    out = [r for r in rows if _key(r) not in seen]
    return out, len(rows) - len(out)


def _num(v, default=-1):
    try:
        return float(v) if v not in (None, '') else default
    except Exception:
        return default


def sort_active(rows):
    """京东API块：围观数降序（缺失沉底，稳定排序）。"""
    return sorted(rows, key=lambda r: _num(r[WATCH_I]), reverse=True)


def sort_closed(rows):
    """京东成交API块：拍卖开始时间日期降序，同日按出价次数降序（缺失沉底）。"""
    def key(r):
        date = str(r[START_I] or '')[:10] or '0000-00-00'
        return (date, _num(r[BID_I]))
    return sorted(rows, key=key, reverse=True)


def _valued_max_row(ws):
    """有值最大行（值非 None/'' 的最大行号）。"""
    m = 0
    for (r, _c), cell in ws._cells.items():
        if cell.value not in (None, ''):
            if r > m:
                m = r
    return m


def first_free_row(ws):
    """首个可用行：空表（A1 为空）返回 1，否则 max_row + 1。

    虚胖检测：openpyxl 的 max_row 含无值空壳单元格，历史删行残留会把
    used-range 撑大。若有值最大行 < max_row，向 stderr 告警并以有值行+1
    为落点，避免中间留下空洞。
    """
    if ws.max_row == 1 and ws.cell(1, 1).value in (None, ''):
        return 1
    vmax = _valued_max_row(ws)
    if vmax and vmax < ws.max_row:
        print(f'WARN: used-range虚胖 max_row={ws.max_row} 有值最大行={vmax}，'
              f'落点取{vmax + 1}', file=sys.stderr)
        return vmax + 1
    return ws.max_row + 1


def append_block(xlsx, sheet, rows, mode, date=None):
    """mode: 'active'（京东API）或 'closed'（京东成交API）。
    去重三段：块内去重 → 排序 → 历史去重（与本 sheet 全部历史行比对，同 key
    一律剔除，只保留新增）。返回 (日期行号, 写入数, 块内去重数, 历史去重数)。
    全被历史命中时仍写日期行（当日心跳，数据行=0）。
    """
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    rows, ndup = dedupe(rows)
    rows = sort_active(rows) if mode == 'active' else sort_closed(rows)
    rows, nhist = dedupe_history(rows, history_keys(ws))
    d = date or datetime.date.today()
    datestr = d.strftime('%Y年%m月%d日') if isinstance(d, (datetime.date, datetime.datetime)) else str(d)
    r0 = first_free_row(ws)
    ws.cell(r0, 1, datestr)
    for i, r in enumerate(rows, r0 + 1):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, v)
    wb.save(xlsx)
    print(f'written: sheet={sheet} 日期行={r0} ({datestr}) 数据行={len(rows)} 去重={ndup} 历史去重={nhist}')
    return r0, len(rows), ndup, nhist


if __name__ == '__main__':
    print('import-only module', file=sys.stderr)
